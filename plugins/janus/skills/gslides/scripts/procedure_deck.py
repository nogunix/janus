#!/usr/bin/env python3
"""Procedure / hands-on decks in a Google Slides brand template, built as PPTX.

Workflow 3 of SKILL.md. The YAML driver (build_gslides.py) creates slides
through the Slides API, which cannot reproduce a corporate template's masters,
divider art and fonts faithfully. For step-by-step procedure decks
(screenshots + numbered steps + copy-paste code) it is more reliable to:

  1. Export the brand template from Google Slides as PPTX, keep exactly four
     slides -- cover, divider, interior (title + body), closing -- and save it
     as template-base.pptx. Do NOT commit it if the template is internal.
  2. Build the deck with python-pptx on those four prototypes (this module).
  3. Upload with conversion:  gws drive files create ... --upload deck.pptx
     (new file) or files update (overwrite) -- see SKILL.md Workflow 3.
  4. QA by exporting the converted file back to PDF from Drive.

Example:

    from procedure_deck import ProcedureDeck, NOTEBG, OKBG, CODEBG
    pd = ProcedureDeck("template/template-base.pptx",
                       shot_dirs=["shots-annot", "shots"],
                       cover_marker="Presentation title")
    pd.cover("手順書（サンプル）\\nハンズオン")
    pd.divider("❶ 脆弱性を確認する", 7, "環境全体のリスクを見渡す")
    pd.two_shots("❶-1 ダッシュボード", "Dashboard", "❶ 7分",
                 "01-dashboard", "01-search", "① …", "② …",
                 note="…", note_fill=OKBG)
    s = pd.one_shot("❸-3 コマンドを実行", "…", "❸ 10分", "b5-terminal",
                    "⑦ 次のコマンドを実行する", code="cat /etc/os-release",
                    code_y=2.05, body_h=1.0)
    pd.save("deck.pptx")

Geometry is in inches on a 10 x 5.625 in (16:9) slide. Layout constants match
the template the module was first built for; pass different fonts through the
constructor for another template. cover_marker is the text the template's
cover title box starts with (its placeholder text), so cover() can find it.

Needs: python-pptx.
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

RED = RGBColor(0xEE, 0x00, 0x00)
TITLE_RED = RGBColor(0xFF, 0x00, 0x00)
DARK = RGBColor(0x15, 0x15, 0x15)
GREY = RGBColor(0x4D, 0x4D, 0x4D)
LTGREY = RGBColor(0x8A, 0x8D, 0x90)
CODEBG = RGBColor(0xF2, 0xF2, 0xF2)   # copy-paste command / YAML boxes
NOTEBG = RGBColor(0xFF, 0xF4, 0xCC)   # ⚠ caveats
OKBG = RGBColor(0xE9, 0xF7, 0xDF)     # "what this shows" / confirmation
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xC7, 0xC7, 0xC7)


class ProcedureDeck:
    def __init__(self, template, shot_dirs=(), font="M PLUS Rounded 1c",
                 title_font="M PLUS 1p", mono="Roboto Mono", cover_marker=None):
        self.prs = Presentation(template)
        slides = list(self.prs.slides)
        if len(slides) != 4:
            raise ValueError("template must hold exactly 4 slides: cover, divider, interior, closing")
        self.cover_slide, divider_proto, interior_proto, _closing = slides
        self._divider_layout = divider_proto.slide_layout
        self._interior_layout = interior_proto.slide_layout
        self.shot_dirs = [Path(d) for d in shot_dirs]
        self.font, self.title_font, self.mono = font, title_font, mono
        self.cover_marker = cover_marker

    # ------------------------------------------------------------ primitives
    def shot(self, name):
        """First existing <dir>/<name>.png across shot_dirs (annotated dir first)."""
        for d in self.shot_dirs:
            p = d / f"{name}.png"
            if p.exists():
                return str(p)
        raise FileNotFoundError(f"{name}.png not in {[str(d) for d in self.shot_dirs]}")

    def _font(self, run, size, color=None, bold=False, font=None):
        font = font or self.font
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        run.font.name = font
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):   # East-Asian / complex-script faces too
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", font)

    def _fill(self, tf, body, size, color, bold, font, align, spacing):
        tf.clear()
        for i, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            # a line wrapped in **...** is bold
            is_bold = bold or (line.startswith("**") and line.endswith("**"))
            r = p.add_run()
            r.text = line.strip("*") if is_bold and not bold else line
            self._font(r, size, color, is_bold, font)

    def text(self, slide, x, y, w, h, body, size=11, color=DARK, bold=False, font=None,
             fill=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15,
             shape=MSO_SHAPE.RECTANGLE, margin=0.06):
        box = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            box.fill.background()
        else:
            box.fill.solid()
            box.fill.fore_color.rgb = fill
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Inches(margin))
        self._fill(tf, body, size, color, bold, font, align, spacing)
        return box

    def code(self, slide, x, y, w, h, body, size=9):
        """Monospace copy-paste box (commands, YAML)."""
        return self.text(slide, x, y, w, h, body, size=size, font=self.mono, fill=CODEBG,
                         spacing=1.05, margin=0.12)

    def image(self, slide, name, x, y, w):
        pic = slide.shapes.add_picture(self.shot(name), Inches(x), Inches(y), width=Inches(w))
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(0.75)
        return pic

    def table(self, slide, x, y, w, header, rows, col_w, size=10, row_h=0.32):
        shp = slide.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y),
                                     Inches(w), Inches(row_h * (len(rows) + 1)))
        tbl = shp.table
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
        for r, vals in enumerate([header] + rows):
            tbl.rows[r].height = Inches(row_h)
            for c, v in enumerate(vals):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK if r == 0 else (WHITE if r % 2 else RGBColor(0xF5, 0xF5, 0xF5))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                for side in ("left", "right", "top", "bottom"):
                    setattr(cell, f"margin_{side}", Inches(0.05))
                tf = cell.text_frame
                tf.word_wrap = True
                for i, line in enumerate(str(v).split("\n")):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    run = p.add_run()
                    run.text = line
                    self._font(run, size, WHITE if r == 0 else DARK, bold=(r == 0))
        return shp

    # ------------------------------------------------------------ slide kinds
    @staticmethod
    def _drop_empty_placeholders(slide, keep_idx):
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in keep_idx:
                ph._element.getparent().remove(ph._element)

    @staticmethod
    def _slide_number(slide):
        """Slide-number field in the layout's idx-12 box (python-pptx does not copy it)."""
        ph = next((p for p in slide.placeholders if p.placeholder_format.idx == 12), None)
        if ph is None:
            src = next((p for p in slide.slide_layout.placeholders
                        if p.placeholder_format.idx == 12), None)
            if src is None:
                return
            slide.shapes._spTree.append(copy.deepcopy(src._element))
            ph = next(p for p in slide.placeholders if p.placeholder_format.idx == 12)
        tf = ph.text_frame
        tf.clear()
        para = tf.paragraphs[0]._p
        fld = para.makeelement(qn("a:fld"), {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}",
                                             "type": "slidenum"})
        rpr = fld.makeelement(qn("a:rPr"), {"lang": "ja-JP", "sz": "700"})
        t = fld.makeelement(qn("a:t"), {})
        t.text = "‹#›"
        fld.append(rpr)
        fld.append(t)
        end = para.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(fld)
        else:
            para.append(fld)

    def cover(self, title, size=20, font="M PLUS 1p Medium"):
        """Rewrite the cover's title box (the shape whose text starts with cover_marker)."""
        if not self.cover_marker:
            raise ValueError("pass cover_marker: the text the template's cover title box starts with")
        for sh in self.cover_slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith(self.cover_marker):
                self._fill(sh.text_frame, title, size, WHITE, False, font, PP_ALIGN.LEFT, 1.1)
                sh.top = Inches(2.35)
                sh.height = Inches(0.9)
                return sh
        raise ValueError(f"no cover text box starting with {self.cover_marker!r}")

    def interior(self, title, sub=None, badge=None, title_size=24):
        """Body slide: red title, optional grey subtitle, optional red badge (e.g. '❶ 5分')."""
        s = self.prs.slides.add_slide(self._interior_layout)
        self._drop_empty_placeholders(s, keep_idx={0, 12})
        self._slide_number(s)
        t = s.placeholders[0]
        t.left, t.top, t.width, t.height = Inches(0.43), Inches(0.08), Inches(8.0), Inches(0.5)
        t.text_frame.word_wrap = True
        self._fill(t.text_frame, title, title_size, TITLE_RED, True, self.title_font, PP_ALIGN.LEFT, 1.0)
        if badge:
            self.text(s, 8.55, 0.14, 1.15, 0.34, badge, size=11, color=WHITE, bold=True, fill=RED,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE, margin=0.02)
        if sub:
            self.text(s, 0.43, 0.62, 9.2, 0.32, sub, size=11.5, color=GREY)
        return s

    def divider(self, title, minutes, desc):
        s = self.prs.slides.add_slide(self._divider_layout)
        self._drop_empty_placeholders(s, keep_idx={0, 12})
        self._slide_number(s)
        t = s.placeholders[0]
        t.left, t.top, t.width, t.height = Inches(4.05), Inches(1.34), Inches(5.62), Inches(1.6)
        t.text_frame.word_wrap = True
        self._fill(t.text_frame, title, 25, None, False, self.font, PP_ALIGN.LEFT, 1.1)
        for para in t.text_frame.paragraphs:
            for r in para.runs:
                r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        self.text(s, 4.05, 3.0, 5.4, 0.7, desc, size=12, color=GREY)
        self.text(s, 4.12, 3.8, 1.3, 0.42, f"目安 {minutes}分", size=13, color=WHITE, bold=True,
                  fill=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, margin=0.02)
        return s

    def two_shots(self, title, sub, badge, left, right, cap_l, cap_r, note=None,
                  note_fill=NOTEBG, img_y=1.0):
        """Two screenshots side by side with a caption under each and an optional note bar."""
        s = self.interior(title, sub, badge)
        self.image(s, left, 0.43, img_y, 4.45)
        self.image(s, right, 5.12, img_y, 4.45)
        cy = img_y + 2.55
        self.text(s, 0.43, cy, 4.45, 0.7, cap_l, size=10)
        self.text(s, 5.12, cy, 4.45, 0.7, cap_r, size=10)
        if note:
            self.text(s, 0.43, 4.42, 9.14, 0.42, note, size=10.5, fill=note_fill,
                      anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
        return s

    def one_shot(self, title, sub, badge, name, body, note=None, note_fill=NOTEBG, note_h=1.2,
                 code=None, code_y=None, body_h=2.6):
        """One screenshot left, steps right; optional code box under the steps and note at bottom."""
        s = self.interior(title, sub, badge)
        self.image(s, name, 0.43, 1.0, 5.6)
        self.text(s, 6.25, 1.0, 3.4, body_h, body, size=10.5, spacing=1.3)
        if code:
            self.text(s, 6.25, code_y, 3.4, 0.42, code, size=10, font=self.mono, fill=CODEBG,
                      anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
        if note:
            self.text(s, 6.25, 4.42 - note_h + 0.42, 3.4, note_h, note, size=9.5, fill=note_fill,
                      spacing=1.25, margin=0.1)
        return s

    # ------------------------------------------------------------ finish
    def save(self, out):
        """Move the closing to the end, drop the divider/interior prototypes, renumber parts."""
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        closing = ids[3]
        lst.remove(closing)
        lst.append(closing)
        for proto in (ids[1], ids[2]):
            self.prs.part.drop_rel(proto.get(qn("r:id")))
            lst.remove(proto)
        # same partname-collision guard as deck/decklib.py save()
        self.prs.part.rename_slide_parts([sldId.rId for sldId in lst])
        self.prs.save(out)
        return out

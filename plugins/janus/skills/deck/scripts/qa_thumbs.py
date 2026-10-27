#!/usr/bin/env python3
"""Per-slide visual QA without LibreOffice (macOS).

Some macOS LibreOffice builds drop every CJK glyph when converting
pptx -> pdf: the text is in the PDF (pdftotext extracts it) but nothing
renders, and neither font installs nor cache resets fix it. topdf.sh +
render.sh then cannot be used for QA. macOS QuickLook renders the same pptx
correctly, but only ever produces a thumbnail of slide 1.

So: split the deck into N single-slide decks, QuickLook each one, and tile
the results into a contact sheet.

Previous s*.pptx.png / montage.png in OUTDIR are deleted first, so a reused
directory never mixes a stale run into the contact sheet.

Usage:  python3 qa_thumbs.py DECK.pptx OUTDIR     -> OUTDIR/montage.png
Needs:  python-pptx, qlmanage (macOS), ImageMagick `montage`.
"""

import os
import subprocess
import sys
import tempfile

from pptx import Presentation

RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def keep_only(src, index, dest):
    """Write a copy of `src` containing just slide `index`."""
    prs = Presentation(src)
    xml_slides = prs.slides._sldIdLst
    for i, sld in enumerate(list(xml_slides)):
        if i != index:
            prs.part.drop_rel(sld.get(RID))
            xml_slides.remove(sld)
    prs.save(dest)


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        return 2
    deck, outdir = argv[1], argv[2]
    os.makedirs(outdir, exist_ok=True)
    for f in os.listdir(outdir):
        if (f.startswith("s") and f.endswith(".pptx.png")) or f == "montage.png":
            os.remove(os.path.join(outdir, f))

    total = len(Presentation(deck).slides)
    print(f"{deck}: {total} slides")
    failed = 0
    with tempfile.TemporaryDirectory() as td:
        for i in range(total):
            one = os.path.join(td, f"s{i + 1:02d}.pptx")
            keep_only(deck, i, one)
            subprocess.run(["qlmanage", "-t", "-s", "1600", "-o", outdir, one],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
            if not os.path.exists(os.path.join(outdir, f"s{i + 1:02d}.pptx.png")):
                failed += 1
                print(f"  slide {i + 1:2d}: FAILED")

    pngs = sorted(f for f in os.listdir(outdir) if f.endswith(".pptx.png"))
    if pngs:
        subprocess.run(["montage", *[os.path.join(outdir, p) for p in pngs],
                        "-tile", "3x", "-geometry", "+6+6", "-background", "grey",
                        os.path.join(outdir, "montage.png")], check=False)
        print(f"montage -> {os.path.join(outdir, 'montage.png')}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
